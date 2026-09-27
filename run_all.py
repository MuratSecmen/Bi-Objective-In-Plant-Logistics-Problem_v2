"""Batch driver for the thesis computational campaign.

For every instance (case, n) it runs, in order:
  1. MIP augmented epsilon-constraint sweep (optional per tier),
  2. conversion of the verified MIP front to a reference-fronts JSON,
  3. repeated experiment: NSGA-II A, NSGA-II B and the Solomon heuristic.

Each stage writes a done-flag, so an interrupted campaign resumes where it
stopped. Wall-clock times of all stages go to campaign_timings.csv.
"""
import argparse
import csv
import json
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

REPO = Path(__file__).resolve().parent
BASE_CONFIG = REPO / "inputs_new" / "config.xlsx"

TIERS = {
    "tier1": dict(cases=["case1", "case2", "case3"], sizes=[5, 10],
                  mip=True, mip_time_limit=1800, n_seeds=10, pop=100),
    "tier2": dict(cases=["case1", "case2", "case3"], sizes=[20, 30],
                  mip=False, mip_time_limit=300, n_seeds=10, pop=100),
    "pilot": dict(cases=["case1"], sizes=[10],
                  mip=False, mip_time_limit=60, n_seeds=1, pop=100),
}


def n_generations(n):
    return 100 if n <= 10 else 200


def write_config(path, values):
    wb = openpyxl.load_workbook(BASE_CONFIG)
    ws = wb["config"]
    found = set()
    for row in range(1, ws.max_row + 1):
        key = ws.cell(row, 1).value
        if key in values:
            ws.cell(row, 2).value = values[key]
            found.add(key)
    missing = set(values) - found
    if missing:
        raise KeyError(f"keys not present in config.xlsx: {sorted(missing)}")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def run_stage(cmd, log_path, timings, instance, stage):
    started = time.perf_counter()
    stamp = datetime.now().isoformat(timespec="seconds")
    print(f"[{stamp}] {instance} :: {stage} ...", flush=True)
    with open(log_path, "w", encoding="utf-8") as stream:
        code = subprocess.run(cmd, cwd=REPO, stdout=stream,
                              stderr=subprocess.STDOUT).returncode
    seconds = time.perf_counter() - started
    print(f"    exit={code}  {seconds / 60:.1f} min  log={log_path}", flush=True)
    new_file = not timings.exists()
    with open(timings, "a", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        if new_file:
            writer.writerow(["instance", "stage", "started", "seconds", "exit_code"])
        writer.writerow([instance, stage, stamp, round(seconds, 1), code])
    return code


def mip_front_to_json(mip_dir, case, n, json_path):
    """Collect verified MIP points (feasible sweep iterations) into the
    reference-fronts schema expected by experiments.py."""
    summaries = sorted(mip_dir.rglob("pareto_summary*.xlsx"))
    if not summaries:
        return False
    df = pd.read_excel(summaries[-1])
    if "route_duration_min" not in df.columns:
        return False
    df = df.dropna(subset=["route_duration_min", "total_wait_min"])
    points = [[float(a), float(b)] for a, b in
              zip(df["route_duration_min"], df["total_wait_min"])]
    if not points:
        return False
    json_path.write_text(json.dumps(
        {"product_set_id": case, "num_products": n, "fronts": [points]},
        indent=2), encoding="utf-8")
    return True


def write_meta(path, args):
    """Record what is needed to reproduce the campaign."""
    import platform
    from importlib.metadata import PackageNotFoundError, version

    def pkg(name):
        try:
            return version(name)
        except PackageNotFoundError:
            return None

    try:
        commit = subprocess.run(["git", "rev-parse", "HEAD"], cwd=REPO,
                                capture_output=True, text=True).stdout.strip()
        dirty = bool(subprocess.run(["git", "status", "--porcelain"], cwd=REPO,
                                    capture_output=True, text=True).stdout.strip())
    except OSError:
        commit, dirty = None, None
    meta = dict(
        started=datetime.now().isoformat(timespec="seconds"),
        tier=args.tier, tier_spec=TIERS[args.tier], first_seed=args.seed,
        git_commit=commit, uncommitted_changes=dirty,
        python=sys.version, platform=platform.platform(),
        processor=platform.processor(),
        packages={p: pkg(p) for p in ("gurobipy", "numpy", "pandas", "scipy",
                                      "openpyxl", "matplotlib")},
    )
    path.write_text(json.dumps(meta, indent=2), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--tier", choices=sorted(TIERS), required=True)
    parser.add_argument("--root", type=Path, default=REPO / "results" / "campaign")
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--skip-mip", action="store_true")
    parser.add_argument("--only", nargs="*", default=None,
                        help="restrict to instances such as case1_n10")
    args = parser.parse_args()

    spec = TIERS[args.tier]
    root = args.root
    root.mkdir(parents=True, exist_ok=True)
    timings = root / "campaign_timings.csv"
    write_meta(root / f"campaign_meta_{args.tier}.json", args)

    for case in spec["cases"]:
        for n in spec["sizes"]:
            instance = f"{case}_n{n}"
            if args.only and instance not in args.only:
                continue
            inst_dir = root / instance
            inst_dir.mkdir(parents=True, exist_ok=True)
            gens = 10 if args.tier == "pilot" else n_generations(n)
            common = {
                "product_set_id": case, "num_products": n,
                "auto_visualize": False, "auto_verify": 1,
                "verify_on_fail": "raise",
            }

            ref_json = inst_dir / "mip_reference_fronts.json"
            mip_done = inst_dir / "done_mip.flag"
            if spec["mip"] and not args.skip_mip and not mip_done.exists():
                cfg_path = inst_dir / "config_mip.xlsx"
                write_config(cfg_path, {**common, "run_mode": "multi_objective",
                                        "time_limit_seconds": spec["mip_time_limit"]})
                code = run_stage(
                    [sys.executable, "run_model.py", "--config", str(cfg_path),
                     "--output-dir", str(inst_dir / "mip")],
                    inst_dir / "mip_console.log", timings, instance, "mip")
                if code == 0:
                    mip_done.touch()
            if mip_done.exists() and not ref_json.exists():
                if mip_front_to_json(inst_dir / "mip", case, n, ref_json):
                    print(f"    MIP reference front -> {ref_json}")

            exp_done = inst_dir / "done_experiment.flag"
            if exp_done.exists():
                print(f"{instance} :: experiment already done, skipping")
                continue
            values = dict(common, run_mode="nsga2_new", nsga2_n_workers=8)
            for design in ("nsga2_old", "nsga2_new"):
                values.update({
                    f"{design}_pop_size": spec["pop"],
                    f"{design}_n_gen": gens,
                    f"{design}_time_limit_sec": 7200,
                })
            cfg_path = inst_dir / "config_experiment.xlsx"
            write_config(cfg_path, values)
            methods = (["nsga2_new"] if args.tier == "pilot"
                       else ["nsga2_old", "nsga2_new", "heuristic"])
            cmd = [sys.executable, "run_model.py", "--config", str(cfg_path),
                   "--output-dir", str(inst_dir / "experiment"),
                   "--methods", *methods,
                   "--n-seeds", str(spec["n_seeds"]), "--seed", str(args.seed),
                   "--independent"]
            if ref_json.exists():
                cmd += ["--reference-fronts", str(ref_json)]
            code = run_stage(cmd, inst_dir / "experiment_console.log",
                             timings, instance, "experiment")
            if code in (0, 3):
                exp_done.touch()

    if args.tier == "pilot":
        logs = sorted(root.rglob("convergence_log_*.csv"))
        if logs:
            rows = list(csv.DictReader(open(logs[-1], encoding="utf-8")))
            elapsed = [float(r["elapsed_s"]) for r in rows]
            steps = [b - a for a, b in zip(elapsed, elapsed[1:])]
            tail = steps[len(steps) // 2:] or [elapsed[-1]]
            per_gen = sum(tail) / len(tail)  # steady state, not warm-up
            print(f"\nPILOT: {per_gen:.1f} s per generation (case1, n=10, pop 100)")
            print(f"  one run of 100 generations  ~ {per_gen * 100 / 60:.0f} min")
            print(f"  one instance (2 designs x 10 seeds) ~ "
                  f"{per_gen * 100 * 20 / 3600:.1f} h")


if __name__ == "__main__":
    main()
